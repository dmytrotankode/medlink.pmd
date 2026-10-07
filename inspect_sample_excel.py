import pandas as pd
import openpyxl

path = r"C:\Projects\MedProfit\Індикатори якості.xlsx"
wb = openpyxl.load_workbook(path, read_only=True)
print("Sheet names:", wb.sheetnames)

for sheet in wb.sheetnames:
    df = pd.read_excel(path, sheet_name=sheet, nrows=5)
    print(f"\n--- Sheet: {sheet} ---")
    print("Columns:", list(df.columns))
    print("Head:\n", df.head(3))
