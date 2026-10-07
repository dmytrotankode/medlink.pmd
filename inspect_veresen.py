import pandas as pd
import openpyxl
import sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

path = r"C:\__MEDLINK___\PMG\Вересень 26.xlsx"
wb = openpyxl.load_workbook(path, read_only=True)
for s in wb.sheetnames:
    print(f"Sheet: {s}, max_row: {wb[s].max_row}")

df_head = pd.read_excel(path, sheet_name='Розшифровка', header=None, nrows=6)
for r in range(4):
    print(f"\nRow {r}:")
    for c in range(min(20, df_head.shape[1])):
        val = df_head.iloc[r, c]
        if pd.notna(val):
            print(f"  Col {c}: {val}")
