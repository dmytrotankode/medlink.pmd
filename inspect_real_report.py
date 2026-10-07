import openpyxl
import pandas as pd
import sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

path = r"C:\__MEDLINK___\серпень 2.xlsx"
wb = openpyxl.load_workbook(path, read_only=True)
print("Sheets in 'серпень 2.xlsx':", wb.sheetnames)

for sheet in wb.sheetnames:
    print(f"\n==================== SHEET: {sheet} ====================")
    df = pd.read_excel(path, sheet_name=sheet, nrows=5)
    print("Shape (first rows):", df.shape)
    print("Columns:", list(df.columns)[:15])
    print("Head:\n", df.iloc[:3, :8])
