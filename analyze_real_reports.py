import openpyxl
import pandas as pd
import sys
import os

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

f1 = r"C:\__MEDLINK___\PMG\Вересень 26.xlsx"
f2 = r"C:\__MEDLINK___\PMG\02000334_SF_2026_08_20260910.xlsx"

for path in [f1, f2]:
    fname = os.path.basename(path)
    print(f"\n==================================================================")
    print(f"ANALYZING: {fname}")
    print(f"==================================================================")
    
    # Check Rozshyfrovka row count and columns
    # Let's find the header row first
    wb = openpyxl.load_workbook(path, read_only=True)
    sheet_names = wb.sheetnames
    
    for sname in ['Розшифровка', 'Пацієнти', 'Опис помилок']:
        if sname not in sheet_names:
            continue
        ws = wb[sname]
        total_rows = ws.max_row
        print(f"\n--- Sheet: '{sname}' (Total Rows: {total_rows}) ---")
        
        # Read sample
        header_row_idx = 2 if sname == 'Розшифровка' else 2
        df = pd.read_excel(path, sheet_name=sname, header=None, nrows=10)
        # Find which row has text like 'Звітний рік'
        for r_idx in range(min(5, len(df))):
            row_str = " ".join([str(x) for x in df.iloc[r_idx].values if pd.notna(x)])
            if 'звітний рік' in row_str.lower() or 'тип' in row_str.lower() or 'код' in row_str.lower():
                header_row_idx = r_idx
                break
        
        print(f"  Header identified at row {header_row_idx}:")
        cols = [f"Col {c}: {df.iloc[header_row_idx, c]}" for c in range(min(df.shape[1], 46)) if pd.notna(df.iloc[header_row_idx, c])]
        for c_info in cols[:20]:
            print(f"    {c_info}")
        if len(cols) > 20:
            print(f"    ... and {len(cols)-20} more columns (Total columns: {len(cols)})")

    # Let's check packages and errors in Rozshyfrovka
    try:
        df_full = pd.read_excel(path, sheet_name='Розшифровка', skiprows=header_row_idx)
        print(f"\nData Shape in Розшифровка: {df_full.shape}")
        
        # Find package column and status column
        pkg_col = None
        stat_col = None
        err_col = None
        cost_col = None
        
        for c in df_full.columns:
            c_low = str(c).lower()
            if 'пакет' in c_low and not pkg_col:
                pkg_col = c
            if ('включення' in c_low or 'статус' in c_low) and not stat_col:
                stat_col = c
            if ('помилок' in c_low or 'коментар' in c_low) and not err_col:
                err_col = c
            if any(term in c_low for term in ['вартість', 'тариф', 'сума', 'грн', 'ціна', 'cost', 'price', 'sum']):
                cost_col = c
        
        print(f"Package column detected: '{pkg_col}'")
        print(f"Status column detected: '{stat_col}'")
        print(f"Error column detected: '{err_col}'")
        print(f"Cost/Tariff column detected: '{cost_col}' (IS COST PRESENT? {bool(cost_col)})")
        
        if pkg_col:
            print("\nPackage distribution:")
            print(df_full[pkg_col].value_counts().head(8))
            
        if stat_col:
            print("\nStatus distribution:")
            print(df_full[stat_col].value_counts().head(5))
            
    except Exception as e:
        print("Error analyzing Rozshyfrovka:", e)
